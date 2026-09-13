import openpyxl
import pytest

from app.services.data.sheet_resolvers import sheet_resolver
from app.services.data.spreadsheet_query_engine import resolve_sheet_name_in_wb
from app.services.data.workbook_chat_service import workbook_chat_service


@pytest.mark.parametrize('sheets,mention', [
    (["Hà Nội", "Hạ Nội"], "ha noi"),
    (["HN Chính T8", "HN Chính T9"], "HN Chính"),
])
def test_ambiguous_sheet_names_require_disambiguation(sheets, mention):
    for names in (sheets, list(reversed(sheets))):
        with pytest.raises(ValueError):
            resolve_sheet_name_in_wb(mention, names)
        with pytest.raises(ValueError):
            sheet_resolver.resolve_sheet(f'sheet "{mention}"', names[0], names)
    assert resolve_sheet_name_in_wb(sheets[0], sheets) == sheets[0]


@pytest.mark.asyncio
@pytest.mark.parametrize('mention,code', [('ha noi', 'SHEET_AMBIGUOUS'), ('Không tồn tại', 'SHEET_NOT_FOUND')])
async def test_chat_unresolved_sheet_has_no_actions(tmp_path, mention, code):
    path = tmp_path / 'ambiguous.xlsx'
    wb = openpyxl.Workbook()
    wb.active.title = 'Hà Nội'
    wb.active.append(['Thực lĩnh'])
    wb.active.append([100])
    wb.create_sheet('Hạ Nội').append(['Thực lĩnh'])
    wb.save(path)
    wb.close()
    before = path.read_bytes()
    response = await workbook_chat_service.chat(
        file_path=str(path), message=f'Tổng thực lĩnh trong sheet "{mention}"',
        sheet_name='Hà Nội', conversation_id=str(path),
    )
    assert response['error']['code'] == code
    if code == 'SHEET_AMBIGUOUS':
        assert set(response['error']['candidates']) == {'Hà Nội', 'Hạ Nội'}
    assert response['context']['sheet'] is None
    assert response['actions'] == response['pending_actions'] == []
    assert path.read_bytes() == before


@pytest.mark.asyncio
@pytest.mark.parametrize('condition,rows', [('dưới 10 triệu', [2]), ('trên 10 triệu', [3])])
async def test_highlight_proposal_preserves_workbook(tmp_path, condition, rows):
    path = tmp_path / 'payroll.xlsx'
    wb = openpyxl.Workbook()
    wb.active.title = 'Lương'
    wb.active.append(['Mã NV', 'Thực lĩnh'])
    wb.active.append(['NV001', 9000000])
    wb.active.append(['NV002', 12000000])
    wb.save(path)
    wb.close()
    before = path.read_bytes()
    response = await workbook_chat_service.chat(
        file_path=str(path), message=f'Tô vàng các dòng có thực lĩnh {condition}',
        sheet_name='Lương', conversation_id=str(path),
    )
    assert response['actions'] == []
    assert response['pending_actions'][0]['rows'] == rows
    assert response['pending_actions'][0]['requires_confirmation'] is True
    assert path.read_bytes() == before


def test_numeric_identifiers_are_not_summed(tmp_path):
    from app.services.data.spreadsheet_query_engine import spreadsheet_query_engine
    path = tmp_path / 'ids.xlsx'
    wb = openpyxl.Workbook()
    wb.active.append(['Mã NV', 'Thực lĩnh'])
    wb.active.append([1001, 90])
    wb.active.append([1002, 120])
    wb.save(path)
    wb.close()
    schema = spreadsheet_query_engine.get_sheet_schema(str(path), 'Sheet')
    assert schema['columns'][0]['type'] == 'id_code'
    result = spreadsheet_query_engine.aggregate_column(str(path), 'Sheet', 'Mã NV')
    assert result.get('error')
    assert result.get('value') is None
    assert spreadsheet_query_engine.aggregate_column(str(path), 'Sheet', 'Mã NV', 'count')['value'] == 2
    assert spreadsheet_query_engine.aggregate_column(str(path), 'Sheet', 'Thực lĩnh')['value'] == 210


@pytest.mark.asyncio
@pytest.mark.parametrize('prompt', ['Tìm ô trống', 'Xóa màu đánh dấu'])
async def test_analysis_proposes_changes_without_local_or_google_writes(tmp_path, monkeypatch, prompt):
    from app.services.data.google_sheets_service import google_sheets_service
    path = tmp_path / 'proposals.xlsx'
    wb = openpyxl.Workbook()
    wb.active.append(['Mã NV', 'Thực lĩnh'])
    wb.active.append(['NV001', None])
    wb.active['A2'].fill = openpyxl.styles.PatternFill(fill_type='solid', fgColor='FFFF00')
    wb.save(path)
    wb.close()
    before = path.read_bytes()
    writes = []
    async def token(**kwargs):
        return 'fake-token', None
    async def highlight(**kwargs):
        writes.append(kwargs)
        return {}
    monkeypatch.setattr(google_sheets_service, 'get_valid_access_token', token)
    monkeypatch.setattr(google_sheets_service, 'highlight_cells', highlight)
    response = await workbook_chat_service.analyze_action(
        file_path=str(path), prompt=prompt, sheet_name='Sheet', selected_range='B2:B2',
        conversation_id=str(path), data_source_url='https://docs.google.com/spreadsheets/d/1abcdefghijklmnopqrstuvwxyz0123456789/edit',
    )
    assert writes == []
    assert path.read_bytes() == before
    assert response['actions'] == []
    assert response['pending_actions'][0]['requires_confirmation'] is True
    assert response['google_sync']['google_sync_attempted'] is False

@pytest.mark.asyncio
async def test_multi_sheet_analysis_preserves_pending_actions(tmp_path):
    path = tmp_path / 'multi.xlsx'
    wb = openpyxl.Workbook()
    for ws in [wb.active, wb.create_sheet('Second')]:
        ws.append(['Name', 'Salary'])
        ws.append(['Employee', None])
    wb.save(path)
    wb.close()
    response = await workbook_chat_service.analyze_action(
        file_path=str(path), prompt='Tìm ô trống', sheet_name='Sheet',
        scope={'type': 'workbook'}, conversation_id=str(path),
    )
    assert {item['sheet'] for item in response['pending_actions']} == {'Sheet', 'Second'}
    assert response['actions'] == []

def test_case_insensitive_exact_sheet_wins_over_accent_fallback():
    from app.services.data.sheet_resolvers import SheetResolver
    resolved, _ = SheetResolver.resolve_sheet('tổng lương sheet HN CHINH T8', 'HN Chính T8', ['HN Chính T8', 'HN Chinh T8'])
    assert resolved == 'HN Chinh T8'
