import io

import openpyxl
import pandas as pd

from app.services.data.analysis_contracts import AnalysisScope
from app.services.data.data_quality_service import scan_quality


def quality_workbook(path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Tên", "Doanh thu", "Ngày tạo", "Hỗn hợp", "Ghi chú", "Công thức"])
    sheet.append(["A", 10, "2026-01-01", 1, None, "#DIV/0!"])
    sheet.append(["B", 10, "không-phải-ngày", "text", "ok", 2])
    sheet.append(["C", 10, "2026-03-01", 2, "ok", 3])
    sheet.append(["C", 10, "2026-03-01", 2, "ok", 3])
    sheet.append([" D ", 1000, "2026-04-01", 3, "ok", 4])
    clean = workbook.create_sheet("Clean")
    clean.append(["Mã", "Giá trị"])
    clean.append(["A", 1])
    clean.append(["B", 2])
    workbook.save(path)
    workbook.close()


def test_quality_scan_finds_supported_issues_with_bounded_evidence(tmp_path):
    path = tmp_path / "quality.xlsx"
    quality_workbook(path)

    result = scan_quality(str(path), AnalysisScope("sheet", ("Data",)))
    issue_types = {issue["type"] for issue in result["issues"]}

    assert {
        "missing_values",
        "duplicate_rows",
        "outliers",
        "invalid_dates",
        "mixed_types",
        "formula_errors",
        "whitespace",
    } <= issue_types
    assert result["summary"]["issue_count"] == len(result["issues"])
    assert result["summary"]["sheets_scanned"] == 1
    assert result["checks"]["formula_errors"]["supported"] is True
    for issue in result["issues"]:
        assert issue["id"].startswith("dq_")
        assert issue["sheet"] == "Data"
        assert issue["affected_count"] >= 1
        assert len(issue["cells"]) <= 200
        assert issue["method"]
        assert issue["recommendation"]
        assert "sample_rows" not in issue
        assert "preview_rows" not in issue

    missing = next(
        issue
        for issue in result["issues"]
        if issue["type"] == "missing_values" and issue["column"] == "Ghi chú"
    )
    invalid = next(issue for issue in result["issues"] if issue["type"] == "invalid_dates")
    formula = next(issue for issue in result["issues"] if issue["type"] == "formula_errors")
    assert "E2" in missing["cells"]
    assert all(
        "F2" not in issue["cells"]
        for issue in result["issues"]
        if issue["type"] == "missing_values"
    )
    assert "C3" in invalid["cells"]
    assert "F2" in formula["cells"]


def test_quality_scan_is_stable_and_honors_selected_sheets(tmp_path):
    path = tmp_path / "quality.xlsm"
    quality_workbook(path)
    scope = AnalysisScope("sheets", ("Clean", "Data"))

    first = scan_quality(str(path), scope)
    second = scan_quality(str(path), scope)

    assert [issue["id"] for issue in first["issues"]] == [issue["id"] for issue in second["issues"]]
    assert first == second
    assert first["summary"]["sheets_scanned"] == 2
    assert {issue["sheet"] for issue in first["issues"]} == {"Data"}


def test_csv_reports_formula_check_as_unsupported_without_inventing_issue(tmp_path):
    path = tmp_path / "quality.csv"
    pd.DataFrame(
        {
            "Tên": ["A", "B", "B"],
            "Giá trị": [1, None, None],
        }
    ).to_csv(path, index=False)

    result = scan_quality(str(path), AnalysisScope("workbook", ("quality",)))

    assert result["checks"]["formula_errors"] == {
        "supported": False,
        "reason": "CSV không lưu công thức Excel để kiểm tra lỗi.",
    }
    assert all(issue["type"] != "formula_errors" for issue in result["issues"])
    assert all(issue["sheet"] == "quality" for issue in result["issues"])


def test_clean_sheet_has_no_false_positive(tmp_path):
    path = tmp_path / "clean.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Clean"
    sheet.append(["Mã", "Giá trị"])
    for index in range(1, 6):
        sheet.append([f"M{index}", index])
    workbook.save(path)
    workbook.close()

    result = scan_quality(str(path), AnalysisScope("sheet", ("Clean",)))

    assert result["issues"] == []
    assert result["summary"]["high_count"] == 0
