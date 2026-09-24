import io

import openpyxl
import pytest

from app.services.data.route_enrichment import (
    RouteMeasurement,
    Waypoint,
    enrich_workbook,
    load_reference_points,
    parse_route_stops,
)


def workbook_bytes():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Tuyen"
    sheet.append(["STT", "Kho đi", "Lộ trình", "km 1 chiều\nSPX chốt", "Link ggmap", "km 1 chiều\nDP đo"])
    sheet.append([1, "BN Mega SOC", "Hub A -> Hub B -> BN Mega SOC", 20, None, None])
    sheet.append([2, "Hub A", "Hub B", 12, "https://www.google.com/maps/dir/old", 13.2])
    sheet.append([3, "Unknown Seller", "Hub B", 16, None, None])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.parametrize("origin,route,expected", [
    ("A", "B → C", ["A", "B", "C"]),
    ("A", "A → B → C", ["A", "B", "C"]),
    ("A", "B → C → A", ["A", "B", "C", "A"]),
    ("A", "B → A → C", ["A", "B", "A", "C"]),
    ("A", "A → A → B → C", ["A", "B", "C"]),
    ("A", "B → B → C → A", ["A", "B", "C", "A"]),
    ("A", "B → C → A → A", ["A", "B", "C", "A"]),
    ("BN Mega SOC", "22-BNH Bac Ninh 5 Hub → 22-BNH Bac Ninh 4 Hub → BN Mega SOC",
     ["BN Mega SOC", "22-BNH Bac Ninh 5 Hub", "22-BNH Bac Ninh 4 Hub", "BN Mega SOC"]),
    ("HN2 SOC", "Hub A → Hub B → HN2 SOC", ["HN2 SOC", "Hub A", "Hub B", "HN2 SOC"]),
    ("BN Mega SOC", "BN Mega SOC → Hub A → Hub B → BN Mega SOC",
     ["BN Mega SOC", "Hub A", "Hub B", "BN Mega SOC"]),
    ("A", "B\nC\nA", ["A", "B", "C", "A"]),
])
def test_parse_route_stops_preserves_full_ordered_route(origin, route, expected):
    assert parse_route_stops(origin, route) == expected


def test_reference_csv_requires_precise_coordinates_and_detects_conflicts():
    content = "Tên HUB,Vĩ độ,Kinh độ\nBN Mega SOC,21.0738,105.97838\nHub A,21.12,106.01\nHub A,21.13,106.02\nHub B,21.18,106.05\n".encode()
    points = load_reference_points(content, "hubs.csv")
    assert points["bn mega soc"].latitude == 21.0738
    assert points["hub a"] is None
    assert points["hub b"].longitude == 106.05


def test_looker_chart_export_uses_hub_seller_and_google_maps_pin():
    content = (
        "Tỉnh,Hub seller,Định vị\n"
        "Điện Biên,22-DBN Dien Bien Hub,https://www.google.com/maps/place/21.384945%2C+103.032207\n"
        'Bắc Ninh,22-BNH Bac Ninh 5 Hub,"https://www.google.com/maps/place/21.10234,+106.01234"\n'
    ).encode()
    points = load_reference_points(content, "looker-hubs.csv")
    assert points["22-dbn dien bien hub"] == Waypoint("22-DBN Dien Bien Hub", 21.384945, 103.032207)
    assert points["22-bnh bac ninh 5 hub"].longitude == 106.01234


def test_looker_chart_export_rejects_map_viewport_and_invalid_pin():
    content = (
        "Hub seller,Định vị\n"
        'Hub A,"https://www.google.com/maps/@21.123,105.456,12z"\n'
        'Hub B,"https://example.com/maps/place/21.123,105.456"\n'
    ).encode()
    with pytest.raises(ValueError, match="tọa độ hợp lệ"):
        load_reference_points(content, "looker-hubs.csv")


@pytest.mark.asyncio
async def test_enrichment_measures_full_route_and_preserves_old_values_in_audit():
    source = workbook_bytes()
    points = load_reference_points(
        "name,lat,lng\nBN Mega SOC,21.0738,105.97838\nHub A,21.12,106.01\nHub B,21.18,106.05\n".encode(),
        "hubs.csv",
    )
    calls = []

    async def measure(waypoints):
        calls.append([point.name for point in waypoints])
        return RouteMeasurement(31567, (10000, 10000, 11567)) if len(waypoints) == 4 else RouteMeasurement(14000, (14000,))

    result, summary = await enrich_workbook(source, points, measure, "https://docs.google.com/spreadsheets/d/source")
    assert source != result
    assert calls[0] == ["BN Mega SOC", "Hub A", "Hub B", "BN Mega SOC"]
    workbook = openpyxl.load_workbook(io.BytesIO(result))
    sheet = workbook["Tuyen"]
    assert sheet["F2"].value == 31.6
    assert "origin=" in sheet["E2"].value
    assert "destination=" in sheet["E2"].value
    assert sheet["F3"].value is not None
    assert sheet["E3"].value != "https://www.google.com/maps/dir/old"
    assert sheet["F4"].value is None
    assert summary.total == 3
    assert summary.completed == 2
    assert summary.missing_waypoint == 1
    assert summary.round_trips == 1
    audit = workbook["Doi chieu lo trinh"]
    assert audit["A1"].value == "Dòng Excel"
    assert audit["V3"].value == "https://www.google.com/maps/dir/old"
    assert audit["U3"].value == 13.2
    assert "Unknown Seller" in audit["S4"].value
    assert audit["R2"].value == "Cần kiểm tra - KM chênh lệch lớn"


@pytest.mark.asyncio
async def test_identical_routes_are_measured_once():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Kho đi", "Lộ trình", "Link ggmap", "km 1 chiều DP đo"])
    sheet.append(["Hub A", "Hub B", None, None])
    sheet.append(["Hub A", "Hub B", None, None])
    raw = io.BytesIO()
    workbook.save(raw)
    points = load_reference_points(b"name,lat,lng\nHub A,21.12,106.01\nHub B,21.18,106.05\n", "hubs.csv")
    calls = 0

    async def measure(_):
        nonlocal calls
        calls += 1
        return RouteMeasurement(12345, (12345,))

    result, summary = await enrich_workbook(raw.getvalue(), points, measure, "https://example.com/hubs.csv")
    assert summary.completed == 2
    assert calls == 1
    output = openpyxl.load_workbook(io.BytesIO(result))
    assert output.active["D2"].value == output.active["D3"].value == 12.3


@pytest.mark.asyncio
async def test_export_never_marks_round_trip_ok_if_google_omits_return_leg():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Kho đi", "Lộ trình", "KM 1 chiều SPX chốt", "Link ggmap", "km 1 chiều DP đo"])
    sheet.append(["A", "B → C → A", 40, "https://www.google.com/maps/dir/old", 38])
    raw = io.BytesIO()
    workbook.save(raw)
    points = load_reference_points(b"name,lat,lng\nA,21.0,105.0\nB,21.1,105.1\nC,21.2,105.2\n", "hubs.csv")

    async def missing_return_leg(_):
        return RouteMeasurement(25000, (10000, 15000))

    result, summary = await enrich_workbook(raw.getvalue(), points, missing_return_leg, "https://example.com/hubs.csv")
    output = openpyxl.load_workbook(io.BytesIO(result))
    assert summary.total == 1
    assert summary.completed == 0
    assert summary.maps_failed == 1
    assert summary.validation_failed == 1
    assert summary.round_trips == 1
    assert output.active["D2"].value is None
    assert output.active["E2"].value is None
    audit = output["Doi chieu lo trinh"]
    assert audit["J2"].value == 4
    assert audit["K2"].value == 3
    assert audit["L2"].value == "3/4 - FAIL"
    assert audit["H2"].value == "A"
    assert audit["R2"].value == "Lỗi - Google Maps không tạo được tuyến"
    assert audit["U2"].value == 38


@pytest.mark.asyncio
async def test_resolved_waypoint_names_are_checked_at_each_index():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Kho đi", "Lộ trình", "Link ggmap", "km 1 chiều DP đo"])
    sheet.append(["A", "B → C → A", None, None])
    raw = io.BytesIO()
    workbook.save(raw)
    # A malformed reference mapping must not silently substitute B and C.
    points = {
        "a": Waypoint("A", 21.0, 105.0),
        "b": Waypoint("C", 21.2, 105.2),
        "c": Waypoint("B", 21.1, 105.1),
    }

    async def measure(_):
        raise AssertionError("Tuyến sai thứ tự không được gửi Google Routes")

    result, summary = await enrich_workbook(raw.getvalue(), points, measure, "https://example.com/hubs.csv")
    audit = openpyxl.load_workbook(io.BytesIO(result))["Doi chieu lo trinh"]
    assert summary.completed == 0
    assert summary.missing_waypoint == 1
    assert audit["R2"].value == "Lỗi - Không xác định được waypoint"
    assert "#2 B" in audit["S2"].value
