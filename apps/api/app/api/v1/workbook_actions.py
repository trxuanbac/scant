"""Authenticated, version-checked ledger for local spreadsheet display layers."""
import hashlib
import io
import json
from datetime import datetime, timezone
from typing import Literal, Optional

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, Query
from pydantic import BaseModel, Field, ValidationError, ConfigDict
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession
from app.api.deps import get_current_user
from app.core.database import get_db
from app.models.entities import User, AuditLog
from app.models.workbook_action import WorkbookAction
from app.services.data.analysis_source import resolve_analysis_source
from app.services.data.analysis_session_service import owned_session
from app.services.data.data_access import validated_highlight

router = APIRouter(prefix='/data/workbook-actions', tags=['data'])


class LayerSpec(BaseModel):
    model_config = ConfigDict(extra='forbid')
    type: Literal['HIGHLIGHT_CELLS', 'CLEAR_HIGHLIGHTS']
    sheet: str = Field(min_length=1, max_length=31)
    cells: list[str] = Field(default_factory=list, max_length=50000)
    color: str = '#FFFF00'
    label: str = Field(default='', max_length=250)


async def source_bytes(db, user, file, file_id, data_source_url):
    return await resolve_analysis_source(
        db, user, file=file, file_id=file_id, data_source_url=data_source_url
    )


def key_hash(value):
    return hashlib.sha256(value.encode()).hexdigest()


async def lock_owner(db, user):
    # Works for SQLite's single writer and PostgreSQL row locks.
    await db.execute(update(User).where(User.id == user.id).values(id=User.id))


async def records(db, user_id, source_key, source_hash):
    return list((await db.scalars(select(WorkbookAction).where(
        WorkbookAction.user_id == user_id, WorkbookAction.source_key == source_key,
        WorkbookAction.source_hash == source_hash,
    ).order_by(WorkbookAction.applied_at, WorkbookAction.created_at, WorkbookAction.id))).all())


def revision(items):
    return key_hash(json.dumps(sorted((item.id, item.status) for item in items if item.status in {'applied', 'undone'})))


def serialize(item):
    return {'id': item.id, 'status': item.status, 'source_hash': item.source_hash,
            'analysis_session_id': item.analysis_session_id,
            'action': item.payload_json, 'preview': item.preview_json,
            'created_at': item.created_at.isoformat(), 'applied_at': item.applied_at.isoformat() if item.applied_at else None}


def audit(db, user, item, event):
    db.add(AuditLog(user_id=user.id, action=event, resource_type='workbook_layer', resource_id=item.id,
                    details_json={'source_hash': item.source_hash, 'sheet': item.payload_json['sheet'],
                                  'cell_count': len(item.payload_json.get('cells', [])), 'target': 'scant_display'}))


@router.post('/preview')
async def preview_action(source_key: str = Form(..., min_length=1, max_length=1000), action: str = Form(...),
                         file: Optional[UploadFile] = File(None), file_id: Optional[str] = Form(None),
                         data_source_url: Optional[str] = Form(None), user: User = Depends(get_current_user),
                         analysis_session_id: Optional[str] = Form(None), db: AsyncSession = Depends(get_db)):
    try:
        spec = LayerSpec.model_validate_json(action)
    except ValidationError:
        raise HTTPException(422, 'Hành động không hợp lệ hoặc chưa được hỗ trợ.')
    if spec.type == 'HIGHLIGHT_CELLS':
        spec.cells, normalized_color = validated_highlight(json.dumps(spec.cells), spec.color)
        spec.color = "#" + normalized_color[-6:]
    source = await source_bytes(db, user, file, file_id, data_source_url)
    content = source.content
    analysis_session = None
    if analysis_session_id:
        analysis_session = await owned_session(db, user.id, analysis_session_id)
        if analysis_session.source_version != source.source_version.version:
            raise HTTPException(409, 'Phiên phân tích thuộc phiên bản workbook khác.')
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            if spec.sheet not in wb.sheetnames:
                raise HTTPException(422, 'Sheet không tồn tại trong workbook.')
            ws = wb[spec.sheet]
            sample = []
            for address in spec.cells[:200]:
                cell = ws[address]
                sample.append({'cell': address, 'value': cell.value if isinstance(cell.value, (str, int, float, bool, type(None))) else str(cell.value),
                               'after_color': spec.color})
        finally:
            wb.close()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'Lịch sử bền vững hiện hỗ trợ workbook XLSX/XLSM hợp lệ.')
    await lock_owner(db, user)
    digest = hashlib.sha256(content).hexdigest()
    items = await records(db, user.id, key_hash(source_key), digest)
    item = WorkbookAction(user_id=user.id, source_key=key_hash(source_key), source_hash=digest,
                          analysis_session_id=analysis_session.id if analysis_session else None,
                          base_revision=revision(items), payload_json=spec.model_dump(), preview_json=sample)
    db.add(item)
    await db.flush()
    audit(db, user, item, 'WORKBOOK_LAYER_PROPOSED')
    return serialize(item)


@router.get('/history')
async def history(source_key: str = Query(..., max_length=1000), source_hash: str = Query(..., pattern=r'^[a-f0-9]{64}$'),
                  user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    items = await records(db, user.id, key_hash(source_key), source_hash)
    layers = []
    for item in items:
        if item.status != 'applied':
            continue
        if item.payload_json['type'] == 'CLEAR_HIGHLIGHTS':
            layers = [layer for layer in layers if layer['sheet'] != item.payload_json['sheet']]
        else:
            layers.append({'id': item.id, **item.payload_json})
    return {'items': [serialize(item) for item in sorted(items, key=lambda item: (item.created_at, item.id), reverse=True)[:100]], 'layers': layers, 'undo_id': next((item.id for item in reversed(items) if item.status == 'applied'), None), 'revision': revision(items), 'has_state': any(item.status in {'applied', 'undone'} for item in items)}


async def owned_action(db, user, action_id):
    item = await db.scalar(select(WorkbookAction).where(WorkbookAction.id == action_id, WorkbookAction.user_id == user.id))
    if item is None:
        raise HTTPException(404, 'Không tìm thấy hành động.')
    return item


@router.post('/{action_id}/confirm')
async def confirm(action_id: str, file: Optional[UploadFile] = File(None), file_id: Optional[str] = Form(None),
                  data_source_url: Optional[str] = Form(None), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    item = await owned_action(db, user, action_id)
    source = await source_bytes(db, user, file, file_id, data_source_url)
    content = source.content
    if hashlib.sha256(content).hexdigest() != item.source_hash:
        raise HTTPException(409, 'Workbook đã thay đổi. Hãy xem trước lại.')
    await lock_owner(db, user)
    await db.refresh(item)
    if item.status == 'applied':
        return serialize(item)
    if item.status != 'pending':
        raise HTTPException(409, 'Hành động không còn chờ xác nhận.')
    items = await records(db, user.id, item.source_key, item.source_hash)
    if revision(items) != item.base_revision:
        raise HTTPException(409, 'Các lớp màu đã thay đổi. Hãy xem trước lại.')
    item.status = 'applied'
    item.applied_at = datetime.now(timezone.utc)
    audit(db, user, item, 'WORKBOOK_LAYER_APPLIED')
    await db.flush()
    return serialize(item)


@router.post('/{action_id}/cancel')
async def cancel(action_id: str, user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    await lock_owner(db, user)
    item = await owned_action(db, user, action_id)
    if item.status == 'cancelled':
        return serialize(item)
    if item.status != 'pending':
        raise HTTPException(409, 'Chỉ hủy được hành động đang chờ.')
    item.status = 'cancelled'
    audit(db, user, item, 'WORKBOOK_LAYER_CANCELLED')
    return serialize(item)


@router.post('/{action_id}/undo')
async def undo(action_id: str, user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    await lock_owner(db, user)
    item = await owned_action(db, user, action_id)
    if item.status == 'undone':
        return serialize(item)
    applied = [entry for entry in await records(db, user.id, item.source_key, item.source_hash) if entry.status == 'applied']
    if not applied or applied[-1].id != item.id:
        raise HTTPException(409, 'Chỉ hoàn tác hành động được áp dụng gần nhất.')
    item.status = 'undone'
    audit(db, user, item, 'WORKBOOK_LAYER_UNDONE')
    return serialize(item)


@router.post('/source')
async def source_version(file: Optional[UploadFile] = File(None), file_id: Optional[str] = Form(None),
                         data_source_url: Optional[str] = Form(None), user: User = Depends(get_current_user),
                         db: AsyncSession = Depends(get_db)):
    source = await source_bytes(db, user, file, file_id, data_source_url)
    return {
        'source_hash': source.source_version.version,
        'source_version': source.source_version.as_dict(),
    }
