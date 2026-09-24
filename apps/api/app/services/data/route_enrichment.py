"""Resolve spreadsheet route rows using verified coordinates and export a copy."""

import asyncio
import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from typing import Awaitable, Callable, Mapping
from urllib.parse import parse_qs, unquote_plus, urlencode, urlparse

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


@dataclass(frozen=True)
class Waypoint:
    name: str
    latitude: float
    longitude: float


@dataclass(frozen=True)
class RouteMeasurement:
    distance_meters: int
    leg_distances_meters: tuple[int, ...]


class RouteValidationError(ValueError):
    """Google returned a route that cannot prove every requested leg exists."""


@dataclass(frozen=True)
class RouteSummary:
    total: int = 0
    completed: int = 0
    missing_waypoint: int = 0
    maps_failed: int = 0
    round_trips: int = 0
    large_difference: int = 0
    missing_links: int = 0
    validation_failed: int = 0

    @property
    def unresolved(self) -> int:
        return self.missing_waypoint + self.maps_failed


def normalize_name(value: object) -> str:
    text = unicodedata.normalize("NFC", str(value or "").strip().casefold())
    return re.sub(r"\s+", " ", text)


def _header_key(value: object) -> str:
    text = unicodedata.normalize("NFKD", normalize_name(value))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]", "", text.replace("đ", "d"))


def parse_route_stops(origin: str, route: str) -> list[str]:
    stops = [part.strip() for part in re.split(r"\s*(?:->|→|\r?\n)\s*", route or "") if part.strip()]
    ordered: list[str] = []
    for point in [origin.strip(), *stops]:
        if not ordered or normalize_name(ordered[-1]) != normalize_name(point):
            ordered.append(point)
    return ordered


def _rows_from_reference(content: bytes, filename: str):
    if filename.lower().endswith(".csv"):
        decoded = content.decode("utf-8-sig")
        sample = decoded[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            reader = csv.reader(io.StringIO(decoded), dialect)
        except csv.Error:
            reader = csv.reader(io.StringIO(decoded))
        yield "CSV", list(reader)
        return
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Bảng tọa độ chỉ hỗ trợ CSV hoặc XLSX.")
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        for sheet in workbook:
            yield sheet.title, list(sheet.values)
    finally:
        workbook.close()


def _find_reference_columns(rows: list[tuple | list]) -> tuple[int, int, int, int] | None:
    name_keys = {"name", "tenhub", "tenkho", "khodi", "hub", "hubseller", "kho", "tendiem", "diem"}
    lat_keys = {"lat", "latitude", "vido", "toadolat"}
    lng_keys = {"lng", "lon", "long", "longitude", "kinhdo", "toadolng"}
    pin_keys = {"dinhvi", "linkdinhvi", "googlemaps", "linkgooglemaps"}
    for index, row in enumerate(rows[:20]):
        keys = [_header_key(cell) for cell in row]
        name = next((i for i, key in enumerate(keys) if key in name_keys), None)
        lat = next((i for i, key in enumerate(keys) if key in lat_keys), None)
        lng = next((i for i, key in enumerate(keys) if key in lng_keys), None)
        if name is not None and lat is not None and lng is not None:
            return index, name, lat, lng
        pin = next((i for i, key in enumerate(keys) if key in pin_keys), None)
        if name is not None and pin is not None:
            return index, name, pin, pin
    return None


def _coordinates_from_pin(value: object) -> tuple[float, float] | None:
    """Accept an explicit Google Maps pin, never a map viewport coordinate."""
    parsed = urlparse(str(value or "").strip())
    host = (parsed.hostname or "").lower()
    if parsed.scheme != "https" or not (host == "google.com" or host.endswith(".google.com")):
        return None
    raw = ""
    if parsed.path.startswith("/maps/place/"):
        raw = unquote_plus(parsed.path.removeprefix("/maps/place/").split("/", 1)[0])
    elif parsed.path.startswith("/maps/search/"):
        raw = unquote_plus(parsed.path.removeprefix("/maps/search/").split("/", 1)[0])
    elif parsed.path.startswith("/maps"):
        raw = parse_qs(parsed.query).get("query", [""])[0]
    match = re.fullmatch(r"\s*([+-]?\d+(?:\.\d+)?)\s*,\s*([+-]?\d+(?:\.\d+)?)\s*", raw)
    if not match:
        return None
    latitude, longitude = (float(value) for value in match.groups())
    return (latitude, longitude) if -90 <= latitude <= 90 and -180 <= longitude <= 180 else None


def load_reference_points(content: bytes, filename: str) -> dict[str, Waypoint | None]:
    points: dict[str, Waypoint | None] = {}
    found_schema = False
    for _, rows in _rows_from_reference(content, filename):
        columns = _find_reference_columns(rows)
        if columns is None:
            continue
        found_schema = True
        header_index, name_col, lat_col, lng_col = columns
        for row in rows[header_index + 1:]:
            if len(row) <= max(name_col, lat_col, lng_col):
                continue
            name = str(row[name_col] or "").strip()
            if not name:
                continue
            if lat_col == lng_col:
                coordinates = _coordinates_from_pin(row[lat_col])
                if coordinates is None:
                    continue
                latitude, longitude = coordinates
            else:
                try:
                    latitude = float(str(row[lat_col]).replace(",", "."))
                    longitude = float(str(row[lng_col]).replace(",", "."))
                except (ValueError, TypeError):
                    continue
            if not (-90 <= latitude <= 90 and -180 <= longitude <= 180):
                continue
            point = Waypoint(name=name, latitude=latitude, longitude=longitude)
            key = normalize_name(name)
            if key in points and points[key] is not None and (
                points[key].latitude != point.latitude or points[key].longitude != point.longitude
            ):
                points[key] = None
            elif key not in points:
                points[key] = point
    if not found_schema:
        raise ValueError("Bảng nguồn cần Tên HUB + Vĩ độ/Kinh độ, hoặc Hub seller + Định vị Google Maps.")
    if not points:
        raise ValueError("Bảng nguồn không có tọa độ hợp lệ.")
    return points


def _find_route_columns(sheet) -> tuple[int, dict[str, int]] | None:
    wanted = {
        "origin": {"khodi"},
        "route": {"lotrinh"},
        "link": {"linkggmap", "linkgooglemaps", "linkmaps"},
        "distance": {"km1chieudpdo", "kmdpdo"},
        "old_distance": {"km1chieuspxchot", "kmspxchot"},
    }
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20)):
        keys = [_header_key(cell.value) for cell in row]
        positions = {field: next((i + 1 for i, key in enumerate(keys) if key in variants), None)
                     for field, variants in wanted.items()}
        if all(positions[field] for field in ("origin", "route", "link", "distance")):
            return row[0].row, positions
    return None


def _maps_url(waypoints: list[Waypoint]) -> str:
    coordinate = lambda point: f"{point.latitude:.7f},{point.longitude:.7f}"
    params = {
        "api": "1",
        "origin": coordinate(waypoints[0]),
        "destination": coordinate(waypoints[-1]),
        "travelmode": "driving",
    }
    if len(waypoints) > 2:
        params["waypoints"] = "|".join(coordinate(point) for point in waypoints[1:-1])
    return "https://www.google.com/maps/dir/?" + urlencode(params)


def _safe_audit_text(value: object) -> str:
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _numeric(value: object) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return None


def _validate_maps_url(link: str, waypoints: list[Waypoint]) -> None:
    if len(link) > 2048:
        raise RouteValidationError("Link Google Maps vượt giới hạn độ dài, không bảo đảm đủ waypoint.")
    query = parse_qs(urlparse(link).query)
    actual = [*query.get("origin", []), *(query.get("waypoints", [""])[0].split("|") if "waypoints" in query else []),
              *query.get("destination", [])]
    expected = [f"{point.latitude:.7f},{point.longitude:.7f}" for point in waypoints]
    if actual != expected:
        raise RouteValidationError("Link Google Maps không giữ nguyên số lượng hoặc thứ tự waypoint.")


def _validate_measurement(measured: RouteMeasurement, waypoints: list[Waypoint]) -> None:
    if not isinstance(measured, RouteMeasurement):
        raise RouteValidationError("Kết quả Google Routes thiếu thông tin từng chặng.")
    if len(measured.leg_distances_meters) != len(waypoints) - 1:
        raise RouteValidationError("Google Routes trả về thiếu chặng của tuyến.")
    if measured.distance_meters < 0 or any(distance < 0 for distance in measured.leg_distances_meters):
        raise RouteValidationError("Google Routes trả về quãng đường không hợp lệ.")
    if abs(sum(measured.leg_distances_meters) - measured.distance_meters) > 100:
        raise RouteValidationError("Tổng km Google Routes không khớp các chặng.")


async def enrich_workbook(
    content: bytes,
    points: Mapping[str, Waypoint | None],
    measure: Callable[[list[Waypoint]], Awaitable[RouteMeasurement]],
    source_url: str,
) -> tuple[bytes, RouteSummary]:
    workbook = openpyxl.load_workbook(io.BytesIO(content))
    audit_name = "Doi chieu lo trinh"
    summary_name = "Tong hop lo trinh"
    if audit_name in workbook.sheetnames or summary_name in workbook.sheetnames:
        raise ValueError("File đã có sheet đối chiếu; hãy dùng file gốc để tránh ghi đè lịch sử.")
    audit = workbook.create_sheet(audit_name)
    audit.append([
        "Dòng Excel", "Sheet", "Kho đi", "Lộ trình gốc", "Lộ trình chuẩn hóa", "Origin",
        "Waypoints", "Destination", "Round Trip", "Số waypoint kỳ vọng", "Số waypoint thực tế",
        "So sánh waypoint", "KM cũ", "KM Google Maps", "Chênh lệch KM", "Chênh lệch %",
        "Link Google Maps", "Trạng thái", "Waypoint lỗi", "Ghi chú", "KM DP ban đầu",
        "Link gốc", "Nguồn tọa độ",
    ])
    total = completed = missing_waypoint = maps_failed = round_trips = large_difference = validation_failed = 0
    found_route_sheet = False
    pending: list[dict] = []
    audit_entries: list[tuple[int, int, list]] = []

    def append_entry(sheet_index: int, row_number: int, sheet, names: list[str], route: str,
                     original_link: object, original_km: object, old_km: float | None,
                     actual_count: int, maps_km: float | None, link: str | None,
                     status: str, waypoint_error: str = "", note: str = "") -> None:
        expected_count = len(names)
        round_trip = expected_count > 1 and normalize_name(names[0]) == normalize_name(names[-1])
        difference = round(maps_km - old_km, 1) if maps_km is not None and old_km is not None else None
        difference_percent = round(abs(difference) / old_km * 100, 2) if difference is not None and old_km and old_km > 0 else None
        entry = [
            row_number, sheet.title, _safe_audit_text(names[0] if names else ""), _safe_audit_text(route),
            _safe_audit_text(" → ".join(names)), _safe_audit_text(names[0] if names else ""),
            _safe_audit_text(" → ".join(names[1:-1])), _safe_audit_text(names[-1] if names else ""),
            round_trip, expected_count, actual_count,
            f"{actual_count}/{expected_count} - {'PASS' if actual_count == expected_count and maps_km is not None and link is not None else 'FAIL'}",
            old_km, maps_km, difference, difference_percent, link, status,
            _safe_audit_text(waypoint_error), _safe_audit_text(note), original_km,
            _safe_audit_text(original_link), source_url,
        ]
        audit_entries.append((sheet_index, row_number, entry))

    for sheet in workbook:
        if sheet is audit:
            continue
        detected = _find_route_columns(sheet)
        if detected is None:
            continue
        found_route_sheet = True
        header_row, columns = detected
        sheet_index = workbook.sheetnames.index(sheet.title)
        for row_number in range(header_row + 1, sheet.max_row + 1):
            origin_value = sheet.cell(row_number, columns["origin"]).value
            route_value = sheet.cell(row_number, columns["route"]).value
            origin = origin_value.strip() if isinstance(origin_value, str) else ""
            route = route_value.strip() if isinstance(route_value, str) else ""
            if not origin and not route:
                continue
            total += 1
            link_cell = sheet.cell(row_number, columns["link"])
            distance_cell = sheet.cell(row_number, columns["distance"])
            original_link, original_km = link_cell.value, distance_cell.value
            old_km = _numeric(sheet.cell(row_number, columns["old_distance"]).value) if columns["old_distance"] else None
            # Old G/H values may describe a shortened route. Preserve them in the audit,
            # but never present them as newly verified Google results.
            link_cell.value = None
            distance_cell.value = None
            names = parse_route_stops(origin, route) if origin and route else ([origin] if origin else [])
            is_round_trip = len(names) > 1 and normalize_name(names[0]) == normalize_name(names[-1])
            round_trips += int(is_round_trip)
            if not origin or not route or len(names) < 2 or len(names) > 25:
                missing_waypoint += 1
                reason = "Thiếu Kho đi hoặc Lộ trình." if not origin or not route else "Số điểm dừng không hợp lệ (cần 2–25)."
                append_entry(sheet_index, row_number, sheet, names, route, original_link, original_km,
                             old_km, 0, None, None, "Lỗi - Không xác định được waypoint", reason)
                continue
            missing = [(index + 1, name) for index, name in enumerate(names)
                       if normalize_name(name) not in points
                       or points[normalize_name(name)] is None
                       or normalize_name(points[normalize_name(name)].name) != normalize_name(name)]
            if missing:
                missing_waypoint += 1
                detail = "; ".join(f"#{index} {name}" for index, name in missing)
                append_entry(sheet_index, row_number, sheet, names, route, original_link, original_km,
                             old_km, len(names) - len(missing), None, None,
                             "Lỗi - Không xác định được waypoint", detail)
                continue
            waypoints = [points[normalize_name(name)] for name in names]
            pending.append({"sheet": sheet, "sheet_index": sheet_index, "row": row_number,
                            "link_cell": link_cell, "distance_cell": distance_cell,
                            "names": names, "route": route, "points": waypoints,
                            "original_link": original_link, "original_km": original_km, "old_km": old_km})
    if not found_route_sheet:
        raise ValueError("Không thấy các cột Kho đi, Lộ trình, Link ggmap, km 1 chiều DP đo trong file Excel.")

    def route_key(waypoints: list[Waypoint]):
        return tuple((point.latitude, point.longitude) for point in waypoints)

    unique_routes = {route_key(item["points"]): item["points"] for item in pending}
    distances: dict[tuple, RouteMeasurement | ValueError] = {}
    if unique_routes:
        first_key = next(iter(unique_routes))
        try:
            distances[first_key] = await measure(unique_routes[first_key])
        except ValueError as exc:
            distances[first_key] = exc

        semaphore = asyncio.Semaphore(8)

        async def limited_measure(waypoints):
            async with semaphore:
                return await measure(waypoints)

        remaining = [(key, waypoints) for key, waypoints in unique_routes.items() if key != first_key]
        if remaining:
            measured = await asyncio.gather(*(limited_measure(waypoints) for _, waypoints in remaining), return_exceptions=True)
            for (key, _), result in zip(remaining, measured):
                if isinstance(result, Exception) and not isinstance(result, ValueError):
                    raise result
                distances[key] = result

    for item in pending:
        waypoints = item["points"]
        result = distances[route_key(waypoints)]
        actual_count = len(waypoints)
        try:
            if isinstance(result, ValueError):
                raise result
            _validate_measurement(result, waypoints)
            link = _maps_url(waypoints)
            _validate_maps_url(link, waypoints)
            maps_km = round(result.distance_meters / 1000, 1)
        except ValueError as exc:
            maps_failed += 1
            validation_failed += int(isinstance(exc, RouteValidationError))
            if isinstance(result, RouteMeasurement):
                actual_count = len(result.leg_distances_meters) + 1
            append_entry(item["sheet_index"], item["row"], item["sheet"], item["names"], item["route"],
                         item["original_link"], item["original_km"], item["old_km"], actual_count,
                         None, None, "Lỗi - Google Maps không tạo được tuyến", "", str(exc))
            continue
        item["link_cell"].value = link
        item["distance_cell"].value = maps_km
        old_km = item["old_km"]
        difference_percent = abs(maps_km - old_km) / old_km * 100 if old_km and old_km > 0 else None
        is_large = difference_percent is not None and difference_percent >= 10
        large_difference += int(is_large)
        completed += 1
        status = "Cần kiểm tra - KM chênh lệch lớn" if is_large else "OK - Đã đo đủ toàn bộ tuyến"
        append_entry(item["sheet_index"], item["row"], item["sheet"], item["names"], item["route"],
                     item["original_link"], item["original_km"], old_km, actual_count,
                     maps_km, link, status)

    for _, _, entry in sorted(audit_entries, key=lambda item: (item[0], item[1])):
        audit.append(entry)
        status = entry[17]
        fill = "DCFCE7" if status.startswith("OK") else "FEF3C7" if status.startswith("Cần kiểm tra") else "FEE2E2"
        audit.cell(audit.max_row, 18).fill = PatternFill("solid", fgColor=fill)
        if entry[16]:
            audit.cell(audit.max_row, 17).hyperlink = entry[16]
            audit.cell(audit.max_row, 17).style = "Hyperlink"
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    for cell in audit[1]:
        cell.fill = PatternFill("solid", fgColor="172554")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    audit.row_dimensions[1].height = 34
    for column, width in {"A": 13, "B": 22, "C": 29, "D": 55, "E": 70, "F": 28,
                          "G": 55, "H": 30, "I": 14, "J": 17, "K": 17, "L": 20,
                          "M": 15, "N": 18, "O": 17, "P": 17, "Q": 65, "R": 42,
                          "S": 45, "T": 55, "U": 17, "V": 55, "W": 65}.items():
        audit.column_dimensions[column].width = width
    for row in audit.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    missing_links = sum(1 for _, _, entry in audit_entries if not entry[16])
    summary = RouteSummary(total, completed, missing_waypoint, maps_failed, round_trips,
                           large_difference, missing_links, validation_failed)
    overview = workbook.create_sheet(summary_name)
    overview.append(["Chỉ số", "Số dòng"])
    for label, count in [
        ("Tổng số dòng cần xử lý", summary.total),
        ("Số dòng xử lý thành công", summary.completed),
        ("Số dòng lỗi waypoint", summary.missing_waypoint),
        ("Số dòng Google Maps không tạo được tuyến", summary.maps_failed),
        ("Số dòng round trip", summary.round_trips),
        ("Số dòng KM chênh >= 10%", summary.large_difference),
        ("Số dòng thiếu link", summary.missing_links),
        ("Số dòng validation FAIL", summary.validation_failed),
    ]:
        overview.append([label, count])
    overview.column_dimensions["A"].width = 48
    overview.column_dimensions["B"].width = 17
    overview.freeze_panes = "A2"
    for cell in overview[1]:
        cell.fill = PatternFill("solid", fgColor="172554")
        cell.font = Font(color="FFFFFF", bold=True)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue(), summary
