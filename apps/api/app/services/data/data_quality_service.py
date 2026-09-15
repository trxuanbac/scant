"""Deterministic, coordinate-backed workbook data quality scanning."""

from __future__ import annotations

import datetime as dt
import hashlib
import math
import re
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter, range_boundaries

from app.services.data.analysis_contracts import AnalysisScope


CELL_LIMIT = 200
EXCEL_FORMATS = {".xlsx", ".xlsm"}
DATE_HEADER_PATTERN = re.compile(r"(?:ngày|ngay|date|time|thời gian|thoi gian)", re.IGNORECASE)
EXCEL_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA"}
TYPE_ORDER = {"number": 0, "date": 1, "boolean": 2, "text": 3}
ISSUE_ORDER = {
    "formula_errors": 0,
    "invalid_dates": 1,
    "mixed_types": 2,
    "missing_values": 3,
    "duplicate_rows": 4,
    "outliers": 5,
    "whitespace": 6,
}


def _detect_header_row(raw: pd.DataFrame, max_check: int = 6) -> int:
    best_row, best_score = 0, -1
    for index in range(min(len(raw), max_check)):
        values = raw.iloc[index].dropna().tolist()
        if not values:
            continue
        text_count = sum(isinstance(value, str) and bool(value.strip()) for value in values)
        unique_count = len({str(value).strip() for value in values if str(value).strip()})
        score = text_count * 3 + unique_count * 2 + len(values)
        if unique_count >= 2 and score > best_score:
            best_row, best_score = index, score
    return best_row


def _headers(values: list[Any]) -> list[str]:
    output: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values):
        name = str(value).strip() if value is not None and not pd.isna(value) else f"Cột_{index + 1}"
        if not name:
            name = f"Cột_{index + 1}"
        seen[name] = seen.get(name, 0) + 1
        output.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return output


def _load_table(file_path: str, sheet: str) -> tuple[pd.DataFrame, list[int], int]:
    path = Path(file_path)
    if path.suffix.lower() == ".csv":
        try:
            raw = pd.read_csv(path, header=None)
        except UnicodeDecodeError:
            raw = pd.read_csv(path, header=None, encoding="latin-1")
    else:
        raw = pd.read_excel(path, sheet_name=sheet, header=None)
    if raw.empty:
        return pd.DataFrame(), [], 1
    header_index = _detect_header_row(raw)
    body = raw.iloc[header_index + 1 :].copy()
    body.columns = _headers(raw.iloc[header_index].tolist())
    body["__source_row__"] = [index + 1 for index in body.index]
    data_columns = [column for column in body.columns if column != "__source_row__"]
    body = body.dropna(how="all", subset=data_columns)
    rows = [int(value) for value in body.pop("__source_row__").tolist()]
    return body.reset_index(drop=True), rows, header_index + 1


def _selected_sheets(path: Path, scope: AnalysisScope) -> list[str]:
    if path.suffix.lower() == ".csv":
        return [scope.sheets[0] if scope.sheets else (path.stem or "CSV")]
    return list(scope.sheets)


def _apply_range(
    frame: pd.DataFrame, rows: list[int], scope: AnalysisScope, sheet: str
) -> tuple[pd.DataFrame, list[int]]:
    if scope.mode != "range" or scope.sheets[0] != sheet or not scope.cell_range:
        return frame, rows
    min_col, min_row, max_col, max_row = range_boundaries(scope.cell_range)
    selected_columns = list(frame.columns)[min_col - 1 : max_col]
    positions = [index for index, row in enumerate(rows) if min_row <= row <= max_row]
    return frame.iloc[positions][selected_columns].reset_index(drop=True), [rows[index] for index in positions]


def _family(value: Any) -> str:
    if isinstance(value, (dt.datetime, dt.date, pd.Timestamp)):
        return "date"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)) and not isinstance(value, bool) and not (
        isinstance(value, float) and math.isnan(value)
    ):
        return "number"
    return "text"


def _issue(
    issue_type: str,
    sheet: str,
    *,
    title: str,
    message: str,
    severity: str,
    cells: list[str],
    method: str,
    recommendation: str,
    column: str | None = None,
    ranges: list[str] | None = None,
) -> dict[str, Any]:
    unique_cells = list(dict.fromkeys(cells))
    identity = f"{issue_type}|{sheet}|{column or ''}|{'|'.join(unique_cells)}"
    issue_id = "dq_" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:16]
    bounded = unique_cells[:CELL_LIMIT]
    affected_columns = [column] if column else []
    return {
        "id": issue_id,
        "type": issue_type,
        "severity": severity,
        "title": title,
        "message": message,
        "affected_count": len(unique_cells),
        "affected_rows_count": len(unique_cells),
        "affected_columns": affected_columns,
        "sheet": sheet,
        "column": column,
        "cells": bounded,
        "ranges": list(ranges or bounded),
        "method": method,
        "recommendation": recommendation,
        "suggestion": recommendation,
        "supported": True,
    }


def _table_issues(
    frame: pd.DataFrame,
    rows: list[int],
    sheet: str,
    protected_cells: set[str] | None = None,
) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    protected_cells = protected_cells or set()
    column_names = list(frame.columns)
    for column_index, column in enumerate(column_names, start=1):
        series = frame[column]
        missing_positions = [
            index
            for index, value in enumerate(series)
            if (pd.isna(value) or not str(value).strip())
            and f"{get_column_letter(column_index)}{rows[index]}" not in protected_cells
        ]
        if missing_positions:
            cells = [f"{get_column_letter(column_index)}{rows[index]}" for index in missing_positions]
            pct = len(cells) / max(len(frame), 1) * 100
            issues.append(
                _issue(
                    "missing_values",
                    sheet,
                    column=column,
                    title=f"Thiếu dữ liệu tại cột {column}",
                    message=f"Có {len(cells)} ô trống trong cột '{column}'.",
                    severity="high" if pct > 25 else "medium" if pct > 8 else "low",
                    cells=cells,
                    method="Kiểm tra ô rỗng trên toàn bộ phạm vi đã chọn.",
                    recommendation="Bổ sung dữ liệu hoặc xác nhận quy tắc xử lý giá trị thiếu.",
                )
            )

        non_null = [(index, value) for index, value in enumerate(series) if not pd.isna(value) and str(value).strip()]
        families = sorted({_family(value) for _, value in non_null}, key=lambda item: TYPE_ORDER[item])
        if len(families) > 1:
            cells = [f"{get_column_letter(column_index)}{rows[index]}" for index, _ in non_null]
            issues.append(
                _issue(
                    "mixed_types",
                    sheet,
                    column=column,
                    title=f"Kiểu dữ liệu không đồng nhất tại cột {column}",
                    message=f"Cột '{column}' chứa nhiều kiểu dữ liệu: {', '.join(families)}.",
                    severity="medium",
                    cells=cells,
                    method="Phân loại kiểu giá trị number/date/boolean/text.",
                    recommendation="Chuẩn hóa cột về một kiểu dữ liệu trước khi tính toán.",
                )
            )

        if DATE_HEADER_PATTERN.search(str(column)) and non_null:
            invalid_positions = []
            for index, value in non_null:
                if isinstance(value, (dt.date, dt.datetime, pd.Timestamp)):
                    continue
                parsed = pd.to_datetime(str(value), errors="coerce", format="mixed")
                if pd.isna(parsed):
                    invalid_positions.append(index)
            if invalid_positions:
                cells = [f"{get_column_letter(column_index)}{rows[index]}" for index in invalid_positions]
                issues.append(
                    _issue(
                        "invalid_dates",
                        sheet,
                        column=column,
                        title=f"Ngày không hợp lệ tại cột {column}",
                        message=f"Có {len(cells)} giá trị không chuyển đổi được thành ngày.",
                        severity="medium",
                        cells=cells,
                        method="Chuyển đổi ngày nghiêm ngặt trên cột có tên gợi ý ngày/thời gian.",
                        recommendation="Chuẩn hóa định dạng ngày và sửa các giá trị không hợp lệ.",
                    )
                )

        numeric = pd.to_numeric(series, errors="coerce")
        valid_numeric = numeric.dropna()
        if len(valid_numeric) >= 4 and len(valid_numeric) / max(len(non_null), 1) >= 0.7:
            q1, q3 = float(valid_numeric.quantile(0.25)), float(valid_numeric.quantile(0.75))
            iqr = q3 - q1
            low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
            positions = [index for index, value in enumerate(numeric) if pd.notna(value) and (value < low or value > high)]
            if positions:
                cells = [f"{get_column_letter(column_index)}{rows[index]}" for index in positions]
                issues.append(
                    _issue(
                        "outliers",
                        sheet,
                        column=column,
                        title=f"Giá trị ngoại lai tại cột {column}",
                        message=f"Có {len(cells)} giá trị nằm ngoài ngưỡng IQR.",
                        severity="low",
                        cells=cells,
                        method=f"IQR 1.5× với ngưỡng {low:g} đến {high:g}.",
                        recommendation="Xác nhận các giá trị cực trị trước khi dùng trong báo cáo.",
                    )
                )

        whitespace_positions = [
            index
            for index, value in non_null
            if isinstance(value, str) and (value != value.strip() or "  " in value)
        ]
        if whitespace_positions:
            cells = [f"{get_column_letter(column_index)}{rows[index]}" for index in whitespace_positions]
            issues.append(
                _issue(
                    "whitespace",
                    sheet,
                    column=column,
                    title=f"Khoảng trắng thừa tại cột {column}",
                    message=f"Có {len(cells)} ô chứa khoảng trắng thừa.",
                    severity="low",
                    cells=cells,
                    method="So sánh văn bản gốc với giá trị đã trim và gom khoảng trắng.",
                    recommendation="Chuẩn hóa khoảng trắng trước khi đối chiếu hoặc gom nhóm.",
                )
            )

    if len(frame) > 1:
        duplicated = frame.duplicated(keep=False)
        positions = [index for index, value in enumerate(duplicated) if value]
        if positions:
            cells = [f"A{rows[index]}" for index in positions]
            last_column = get_column_letter(max(len(column_names), 1))
            ranges = [f"A{rows[index]}:{last_column}{rows[index]}" for index in positions]
            issues.append(
                _issue(
                    "duplicate_rows",
                    sheet,
                    title="Phát hiện dòng dữ liệu trùng lặp",
                    message=f"Có {len(positions)} dòng thuộc nhóm trùng lặp hoàn toàn.",
                    severity="high" if len(positions) / len(frame) > 0.1 else "medium",
                    cells=cells,
                    ranges=ranges,
                    method="So sánh toàn bộ giá trị trong từng dòng của phạm vi.",
                    recommendation="Xác nhận và loại bỏ bản ghi trùng để tránh tính lặp.",
                )
            )
    return issues


def _excel_cell_state(
    file_path: str, sheets: list[str], scope: AnalysisScope
) -> tuple[list[dict[str, Any]], dict[str, set[str]]]:
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=False)
    try:
        issues: list[dict[str, Any]] = []
        protected: dict[str, set[str]] = {}
        for sheet in sheets:
            worksheet = workbook[sheet]
            bounds = range_boundaries(scope.cell_range) if scope.mode == "range" and scope.cell_range else None
            cells = []
            protected[sheet] = set()
            for row in worksheet.iter_rows():
                for cell in row:
                    if bounds and not (bounds[0] <= cell.column <= bounds[2] and bounds[1] <= cell.row <= bounds[3]):
                        continue
                    if cell.data_type in {"f", "e"} or (
                        isinstance(cell.value, str) and cell.value.upper() in EXCEL_ERRORS
                    ):
                        protected[sheet].add(cell.coordinate)
                    if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value.upper() in EXCEL_ERRORS):
                        cells.append(cell.coordinate)
            if cells:
                issues.append(
                    _issue(
                        "formula_errors",
                        sheet,
                        title="Ô chứa lỗi công thức Excel",
                        message=f"Có {len(cells)} ô chứa mã lỗi Excel có thể xác minh.",
                        severity="high",
                        cells=cells,
                        method="Đọc kiểu ô lỗi và mã lỗi Excel với openpyxl, không tính lại công thức.",
                        recommendation="Mở các ô lỗi, kiểm tra tham chiếu và công thức nguồn.",
                    )
                )
        return issues, protected
    finally:
        workbook.close()


def scan_quality(file_path: str, scope: AnalysisScope) -> dict[str, Any]:
    path = Path(file_path)
    sheets = _selected_sheets(path, scope)
    formula_supported = path.suffix.lower() in EXCEL_FORMATS
    if formula_supported:
        issues, protected_cells = _excel_cell_state(file_path, sheets, scope)
        formula_check = {"supported": True, "reason": None}
    else:
        issues, protected_cells = [], {}
        formula_check = {
            "supported": False,
            "reason": "CSV không lưu công thức Excel để kiểm tra lỗi."
            if path.suffix.lower() == ".csv"
            else "Định dạng này không hỗ trợ kiểm tra lỗi công thức trực tiếp.",
        }
    for sheet in sheets:
        frame, rows, _header_row = _load_table(file_path, sheet)
        frame, rows = _apply_range(frame, rows, scope, sheet)
        if not frame.empty:
            issues.extend(_table_issues(frame, rows, sheet, protected_cells.get(sheet)))
    issues.sort(key=lambda issue: (sheets.index(issue["sheet"]), ISSUE_ORDER[issue["type"]], issue["column"] or "", issue["id"]))
    return {
        "summary": {
            "issue_count": len(issues),
            "high_count": sum(issue["severity"] == "high" for issue in issues),
            "medium_count": sum(issue["severity"] == "medium" for issue in issues),
            "low_count": sum(issue["severity"] == "low" for issue in issues),
            "sheets_scanned": len(sheets),
        },
        "checks": {"formula_errors": formula_check},
        "issues": issues,
    }
